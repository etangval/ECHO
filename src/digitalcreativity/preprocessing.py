"""Audited source adapters. Originals are read only; outputs contain research data.

Optional dependencies for these adapters are installed with the ``research`` extra.
Times are seconds in a declared source clock, never inferred from integer datetime
storage units (which differ between pandas versions).
"""
from collections import Counter, defaultdict
from pathlib import Path
import json
import numpy as np
import pandas as pd
from .pipeline import digest, new_directory, save_json


def epoch_seconds(values):
    parsed = pd.to_datetime(values, utc=True, errors="coerce")
    ns = np.asarray(parsed, dtype="datetime64[ns]").astype(np.int64)
    result = ns.astype(float)/1e9
    result[ns == np.iinfo(np.int64).min] = np.nan
    return result


def merge_intervals(intervals):
    merged = []
    for a, b in sorted((float(a), float(b)) for a,b in intervals if b > a):
        if merged and a <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a,b])
    return merged


def haversine_m(lat1, lon1, lat2, lon2):
    p1,p2 = np.radians(lat1),np.radians(lat2)
    dp,dl = p2-p1,np.radians(np.asarray(lon2)-lon1)
    x = np.sin(dp/2)**2+np.cos(p1)*np.cos(p2)*np.sin(dl/2)**2
    return 6371008.8*2*np.arcsin(np.sqrt(np.clip(x,0,1)))


def displacement_events(gps, threshold_m=50., max_gap_s=900.):
    """First observed exit from an anchor-centred disk, resetting across gaps.

    Event time is the first confirming fix; the crossing is interval-censored.
    This operational spatial-scale event is not a claim to detect true movement onset.
    """
    if threshold_m <= 0 or max_gap_s <= 0:
        raise ValueError("Positive spatial threshold and observation-gap limit required")
    rows, coverage, diagnostics = [],[],[]
    for entity, frame in gps.groupby("entity_id", sort=True):
        f = frame.sort_values("event_time", kind="stable")
        t,lat,lon = (f[c].to_numpy(float) for c in ["event_time","latitude","longitude"])
        valid = np.isfinite(t)&np.isfinite(lat)&np.isfinite(lon)
        t,lat,lon = t[valid],lat[valid],lon[valid]
        refs = f.loc[valid,"source_row"].to_numpy()
        if not len(t): continue
        split = np.r_[0, np.flatnonzero(np.diff(t)>max_gap_s)+1, len(t)]
        for episode,(begin,end) in enumerate(zip(split[:-1],split[1:])):
            if end-begin < 2: continue
            # Adjacent fixes provide a bounded interval of usable trajectory.
            coverage.append(dict(entity_id=str(entity),start_seconds=t[begin],stop_seconds=float(np.nextafter(t[end-1],np.inf)),observation_id=f"{entity}:{episode}"))
            anchor=begin
            for i in range(begin+1,end):
                distance=float(haversine_m(lat[anchor],lon[anchor],lat[i],lon[i]))
                if distance >= threshold_m:
                    rows.append(dict(source="cats",entity_id=str(entity),event_id=f"cats:{entity}:{refs[i]}:{threshold_m:g}",event_time=float(t[i]),event_type="observed_displacement",time_lower=float(t[i-1]),time_upper=float(t[i]),mark_1=distance,mark_2=float(t[i]-t[anchor]),source_row=str(refs[i]),observation_id=f"{entity}:{episode}"))
                    anchor=i
        if len(t)>1:
            d=haversine_m(lat[:-1],lon[:-1],lat[1:],lon[1:]);dt=np.diff(t)
            linked=(dt>0)&(dt<=max_gap_s)
            diagnostics.append(dict(entity_id=str(entity),valid_fixes=len(t),linked_steps=int(linked.sum()),linked_path_length_m=float(d[linked].sum()),median_link_speed_m_s=float(np.median(d[linked]/dt[linked])) if linked.any() else None,median_fix_interval_s=float(np.median(dt)),gaps_over_limit=int((dt>max_gap_s).sum())))
    return pd.DataFrame(rows),pd.DataFrame(coverage),pd.DataFrame(diagnostics)


def write_prepared(folder, events, coverage, report, global_coverage=None):
    folder=Path(folder)
    events=events.sort_values(["event_time","entity_id","event_id"],kind="stable").reset_index(drop=True)
    if events.empty or events.event_id.duplicated().any() or not np.isfinite(events.event_time).all():
        raise ValueError("Prepared events must be nonempty, unique and finite")
    events.to_parquet(folder/"events.parquet",index=False)
    coverage.to_parquet(folder/"observation_intervals.parquet",index=False)
    # Minimal public-tool input, not intended for committing to the code repository.
    events.rename(columns={"event_time":"time_seconds","entity_id":"category"})[["event_id","time_seconds","category"]].to_csv(folder/"sonification_events.csv",index=False,float_format="%.17g")
    intervals=global_coverage if global_coverage is not None else merge_intervals(coverage[["start_seconds","stop_seconds"]].to_numpy())
    pd.DataFrame(intervals,columns=["start_seconds","stop_seconds"]).to_csv(folder/"sonification_coverage.csv",index=False,float_format="%.17g")
    report.update(prepared_events=len(events),entities=int(events.entity_id.nunique()),first_event_seconds=float(events.event_time.min()),last_event_seconds=float(events.event_time.max()),artifacts={f.name:digest(f) for f in folder.iterdir() if f.is_file()})
    save_json(folder/"preparation.json",report)
    return report


def prepare_retail(path, output):
    from openpyxl import load_workbook
    folder=new_directory(output)
    workbook=load_workbook(path,read_only=True,data_only=True)
    invoices={};counts=Counter()
    for sheet in workbook:
        iterator=sheet.iter_rows(values_only=True);columns=list(next(iterator))
        for row_number,row in enumerate(iterator,2):
            d=dict(zip(columns,row));key=str(d["Invoice"]).strip();counts["product_rows"]+=1
            if key not in invoices:invoices[key]=dict(times=set(),customers=set(),countries=set(),refs=[],positive_goods=False,sheets=set())
            x=invoices[key];x["refs"].append(f"{sheet.title}:{row_number}");x["sheets"].add(sheet.title)
            if d["InvoiceDate"] is not None:x["times"].add(d["InvoiceDate"].isoformat())
            if d["Customer ID"] is not None:x["customers"].add(str(int(d["Customer ID"])))
            if d["Country"]:x["countries"].add(d["Country"])
            if d["Quantity"]>0 and d["Price"]>0:x["positive_goods"]=True
        print(json.dumps({"retail_sheet_read":sheet.title}),flush=True)
    workbook.close()
    rows=[];excluded=[];all_invoice=[]
    origin=pd.Timestamp("2009-12-01T00:00:00Z").timestamp()
    for invoice,x in invoices.items():
        flags=[]
        if len(x["customers"])!=1:flags.append("missing_or_ambiguous_customer")
        if len(x["times"])!=1:flags.append("missing_or_ambiguous_time")
        cancelled=invoice.upper().startswith("C")
        if cancelled:flags.append("cancellation")
        if not cancelled and not invoice.isdigit():flags.append("non_numeric_invoice")
        if not x["positive_goods"]:flags.append("no_positive_priced_purchase_line")
        all_invoice.append(dict(invoice=invoice,cancelled=cancelled,customers=sorted(x["customers"]),times=sorted(x["times"]),source_rows=x["refs"],quality_flags=flags))
        if flags:
            excluded.append(dict(invoice=invoice,reasons=";".join(flags)));counts.update(flags);continue
        t=pd.Timestamp(next(iter(x["times"])),tz="UTC").timestamp()-origin
        rows.append(dict(source="retail",entity_id=next(iter(x["customers"])),event_id="invoice:"+invoice,event_time=t,event_type="purchase_invoice",mark_1=len(x["refs"]),mark_2="|".join(sorted(x["countries"])),source_row="|".join(x["refs"])))
    frame=pd.DataFrame(rows)
    lo,hi=0.,pd.Timestamp("2011-12-10T00:00:00Z").timestamp()-origin
    coverage=pd.DataFrame(dict(entity_id=sorted(frame.entity_id.unique()),start_seconds=lo,stop_seconds=hi))
    pd.DataFrame(excluded).to_csv(folder/"excluded_invoices.csv",index=False)
    (folder/"invoice_audit.jsonl").write_text("\n".join(json.dumps(x,ensure_ascii=False) for x in all_invoice),encoding="utf-8")
    return write_prepared(folder,frame,coverage,dict(source="UCI Online Retail II",source_sha256=digest(path),time_origin="2009-12-01 00:00:00 merchant wall clock",timezone="not encoded; UTC suffix used only as a computational convention; not a claim of UTC measurement",time_precision_seconds=60,rule="One unambiguous, non-cancelled numeric invoice with a positive priced goods line; invoice ID unified across both sheets",quality_counts=dict(counts),invoice_ids=len(invoices),invoice_ids_in_multiple_sheets=sum(len(x['sheets'])>1 for x in invoices.values()),coverage_scope="Calendar span of merchant's supplied log, not an independently measured customer enrolment interval; rates concern purchases in this log only"))


def prepare_taxi(path,output,strict=True):
    import pyarrow.parquet as pq
    folder=new_directory(output);frames=[];relaxed_frames=[];counts=Counter()
    origin=pd.Timestamp("2026-01-01T00:00:00Z").timestamp();duration=31*86400.;offset=0
    needed=["tpep_pickup_datetime","tpep_dropoff_datetime","PULocationID","DOLocationID","trip_distance","fare_amount","total_amount"]
    for batch in pq.ParquetFile(path).iter_batches(batch_size=250000,columns=needed):
        d=batch.to_pandas();t=epoch_seconds(d.tpep_pickup_datetime)-origin;stop=epoch_seconds(d.tpep_dropoff_datetime)-origin
        flags=dict(outside_month=(~np.isfinite(t))|(t<0)|(t>=duration),unknown_zone=~d.PULocationID.between(1,263).to_numpy(),negative_distance=(d.trip_distance<0).to_numpy(),negative_fare=(d.fare_amount<0).to_numpy(),negative_total=(d.total_amount<0).to_numpy(),invalid_duration=(~np.isfinite(stop))|(stop<t)|(stop-t>86400))
        counts["input_rows"]+=len(d);counts.update({k:int(v.sum()) for k,v in flags.items()})
        relaxed=~(flags["outside_month"]|flags["unknown_zone"])
        strict_mask=~np.logical_or.reduce(list(flags.values()))
        frame=pd.DataFrame(dict(source="taxi",entity_id=d.PULocationID.astype(str),event_id=[f"taxi:2026-01:{i}" for i in range(offset,offset+len(d))],event_time=t,event_type="pickup",mark_1=d.trip_distance,mark_2=d.fare_amount,source_row=np.arange(offset,offset+len(d)),quality_flags=[";".join(k for k,v in flags.items() if v[i]) for i in range(len(d))]))
        frames.append(frame.loc[strict_mask if strict else relaxed]);relaxed_frames.append(frame.loc[relaxed]);offset+=len(d)
    events=pd.concat(frames,ignore_index=True);relaxed_events=pd.concat(relaxed_frames,ignore_index=True)
    relaxed_events.to_parquet(folder/"relaxed_events.parquet",index=False)
    coverage=pd.DataFrame(dict(entity_id=sorted(events.entity_id.unique()),start_seconds=0.,stop_seconds=duration))
    return write_prepared(folder,events,coverage,dict(source="NYC TLC Yellow Taxi January 2026",source_sha256=digest(path),time_origin="2026-01-01 00:00:00 America/New_York wall clock",time_precision_seconds=1,strict=strict,quality_counts=dict(counts),relaxed_events=len(relaxed_events),excluded_rows=offset-len(events),rule="One pickup per source row; retain zero-duration/distance trips, exclude out-of-month/unknown-zone and strict negative amount or inconsistent-duration flags",coverage_scope="Monthly submitted taxi log; vendor completeness is not guaranteed"))


def prepare_cats(path,reference_path,output,thresholds=(25.,50.,100.),primary=50.,max_gap=900.):
    folder=new_directory(output)
    d=pd.read_csv(path,dtype=str,keep_default_na=False)
    d['source_row']=np.arange(2,len(d)+2)
    visible=d.visible.str.lower().eq('true')
    d=d.loc[visible].copy();origin=pd.Timestamp('2015-03-01T00:00:00Z').timestamp()
    gps=pd.DataFrame(dict(entity_id=d['individual-local-identifier'],event_time=epoch_seconds(d.timestamp)-origin,latitude=pd.to_numeric(d['location-lat']),longitude=pd.to_numeric(d['location-long']),source_row=d.source_row))
    gps.to_parquet(folder/'visible_gps.parquet',index=False)
    reference=pd.read_csv(reference_path,dtype=str,keep_default_na=False)
    deployments=pd.DataFrame(dict(entity_id=reference['animal-id'],start_seconds=epoch_seconds(reference['deploy-on-date'])-origin,stop_seconds=epoch_seconds(reference['deploy-off-date'])-origin))
    deployments.to_parquet(folder/'deployment_intervals.parquet',index=False)
    sensitivity=[]
    for threshold in thresholds:
        e,c,m=displacement_events(gps,threshold,max_gap)
        e.to_parquet(folder/f'events_{threshold:g}m.parquet',index=False)
        c.to_parquet(folder/f'coverage_{threshold:g}m.parquet',index=False)
        sensitivity.append(dict(threshold_m=threshold,events=len(e),entities=int(e.entity_id.nunique()) if len(e) else 0))
        if threshold==primary:events,coverage,movement=e,c,m
    movement.to_csv(folder/'trajectory_metrics.csv',index=False)
    # Sonification uses the documented tag deployment calendar; statistical
    # exposure uses linked valid GPS intervals instead. Neither aligns cats.
    global_cov=merge_intervals(deployments[['start_seconds','stop_seconds']].to_numpy())
    covered=np.zeros(len(events),bool)
    for a,b in global_cov:covered|=(events.event_time>=a)&(events.event_time<b)
    outside=events.loc[~covered];outside.to_parquet(folder/'events_outside_deployment_calendar.parquet',index=False)
    if len(outside):events=events.loc[covered].copy()
    return write_prepared(folder,events,coverage,dict(source="Movebank Pet Cats Australia",source_sha256=digest(path),reference_sha256=digest(reference_path),time_origin="2015-03-01 00:00:00; Movebank exported clock interpreted as UTC, original strings retained in raw",time_precision_seconds=1,primary_threshold_m=primary,max_gap_seconds=max_gap,threshold_sensitivity=sensitivity,visible_fixes=len(gps),outside_global_deployment_calendar=len(outside),event_rule="First observed displacement >= threshold from the previous event/episode anchor; reset across >15-minute gaps; source fix time retained with crossing interval bounds",coverage_scope="Statistics: consecutive visible-fix episodes with <=15-minute gaps. Audio: union of documented deployment calendars. Channels were not recorded simultaneously throughout; no cat population synchrony inference."),global_cov)


def prepare_hippocampus(path,output):
    import h5py
    folder=new_directory(output)
    with h5py.File(path,'r') as f:
        t=f['units/spike_times'][()];end=f['units/spike_times_index'][()];ids=f['units/id'][()]
        # Fixed behavioral support, supplied independently of spiking/aesthetics.
        position_group=f['processing/behavior/SubjectPosition/SpatialSeries']
        pt=position_group['timestamps'][()];xy=position_group['data'][()]
        lo,hi=float(pt[0]),float(np.nextafter(pt[-1],np.inf))
        unit=np.repeat(ids,np.diff(np.r_[0,end]).astype(int));keep=(t>=lo)&(t<hi)
        all_events=pd.DataFrame(dict(source="hippocampus",entity_id=unit.astype(str),event_id=[f'nwb-spike:{i}' for i in range(len(t))],event_time=t,event_type="sorted_spike",source_row=np.arange(len(t))))
        all_events.to_parquet(folder/'all_session_events.parquet',index=False)
        events=all_events.loc[keep].copy()
        coverage=pd.DataFrame(dict(entity_id=ids.astype(str),start_seconds=lo,stop_seconds=hi))
        pos=pd.DataFrame(dict(event_time=pt,x_cm=xy[:,0],y_cm=xy[:,1]))
        # Speed is a derived mark only; no smoothing or time warp of spikes.
        from scipy.ndimage import gaussian_filter1d
        smooth=gaussian_filter1d(xy,3.,axis=0,mode='nearest')
        dt=np.gradient(pt);pos['speed_cm_s']=np.sqrt(np.sum(np.gradient(smooth,axis=0)**2,axis=1))/dt
        pos.to_parquet(folder/'position_covariates.parquet',index=False)
        for name,prefix in [('trials','intervals/trials'),('sleep_states','processing/behavior/SleepStates')]:
            g=f[prefix]
            table={k:[x.decode() if isinstance(x,bytes) else x for x in o[()]] for k,o in g.items()}
            pd.DataFrame(table).to_parquet(folder/(name+'.parquet'),index=False)
    return write_prepared(folder,events,coverage,dict(source="DANDI 000552",version="0.230630.2304",session="e14_2m3_201121",source_sha256=digest(path),time_origin="NWB session-relative seconds",time_precision_seconds=1/30000,all_session_events=len(t),all_session_units=len(ids),primary_interval=[lo,hi],selection_rule="Complete supplied 2D maze-position interval; independent of firing rate and musical output",coverage_scope="All sorted units in the same acquisition session; unit-specific observability/quality metadata absent",speed_method="Gaussian position smoothing sigma=3 samples (~0.1 s), then time derivative; spikes unchanged",limitations=["Cell types and per-unit anatomical labels absent","CA1 provenance comes from source paper/dataset","No stimulation timestamps; maze task is chosen before terminal optotagging described by the source paper"]),[[lo,hi]])
